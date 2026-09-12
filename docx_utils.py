"""DOCX (Word) text extraction utilities."""

from __future__ import annotations

import io
from typing import Optional

import docx


def extract_text_from_docx(file_bytes: bytes, filename: str) -> tuple[str, Optional[str]]:
    """Extract plain text from a DOCX's raw bytes, including table cells.

    Returns a (text, error) tuple. On success, error is None. On failure
    (corrupted or unreadable file), text is "" and error holds a
    human-readable reason.
    """
    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        return "", f"Could not open DOCX: {exc}"

    try:
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
    except Exception as exc:
        return "", f"Failed to extract text: {exc}"

    text = "\n".join(parts).strip()
    if not text:
        return "", "No extractable text found in this document."

    return text, None
